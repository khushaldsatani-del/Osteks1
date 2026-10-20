import io
import re
from datetime import date

from openpyxl import load_workbook

# Parser for the daily "OSTEKS-Auswertungen-<year>.xlsx" workbook - pure
# Python, no database, so the same function serves the manual upload
# (Production View > Upload) and, later, an automatic daily import.
#
# Workbook layout (one sheet per month, e.g. "Oktober 2026"):
#   - a header row whose column A says "Auswertungen", with one column per
#     day ("01.10." ... "31.10.") - the year is only in the sheet name;
#   - then one row per metric: label in column A ("Umsatz KTL",
#     "Auftragsbestand KTL", ...), one value per day column;
#   - "Summe" / "Schnitt" columns after the days, and a small "Heute"
#     summary block below the metrics - both derived, both ignored.
# Rows are found by their label, never by a fixed row number, so a row
# added or moved in the template still imports. Empty cells are "no value"
# and are skipped; a 0 that is actually in a cell is kept.
#
# Cached formula results are read (data_only=True): the file must have
# been saved by Excel, which is how it arrives.

# Every metric is stored under a stable key derived from its label
# ("Auftragsbestand Schäumen" -> "auftragsbestand_schaeumen").
_UMLAUTS = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "ß": "ss", "²": "2", "³": "3", "Ø": "avg", "ø": "avg"})

# The backlog ("Auftragsbestand") rows the dashboard charts. The total is
# never taken from the file - it is the sum of the three, see
# _check_backlog_total().
BACKLOG_METRICS = ("auftragsbestand_ktl", "auftragsbestand_pulver", "auftragsbestand_schaeumen")
BACKLOG_TOTAL_METRIC = "summe_auftragsbestand"

_DAY_HEADER = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})\.?\s*$")
_YEAR = re.compile(r"(20\d{2})")


def metric_key(label: str) -> str:
    text = label.strip().lower().translate(_UMLAUTS)
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _find_header_row(ws) -> int | None:
    for row in range(1, min(ws.max_row, 30) + 1):
        if str(ws.cell(row, 1).value or "").strip().lower() == "auswertungen":
            return row
    return None


def _day_columns(ws, header_row: int, year: int) -> dict[int, date]:
    columns = {}
    for col in range(2, ws.max_column + 1):
        value = ws.cell(header_row, col).value
        if hasattr(value, "year") and hasattr(value, "month"):  # a real date cell
            columns[col] = date(value.year, value.month, value.day)
            continue
        match = _DAY_HEADER.match(str(value or ""))
        if not match:
            continue
        try:
            columns[col] = date(year, int(match.group(2)), int(match.group(1)))
        except ValueError:
            pass  # e.g. "31.02." in a template - not a real day
    return columns


def _is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def parse_workbook(content: bytes) -> dict:
    """Read every monthly sheet. Returns
    {"values": {(iso_date, metric): float}, "metrics": {metric: label},
     "months": [...per-sheet summary...], "warnings": [...]}."""
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    values: dict[tuple[str, str], float] = {}
    metrics: dict[str, dict] = {}
    months: list[dict] = []
    warnings: list[dict] = []

    for ws in wb.worksheets:
        year_match = _YEAR.search(ws.title)
        header_row = _find_header_row(ws)
        if not year_match or header_row is None:
            # e.g. "Nachpflege-Übersicht" - not a monthly sheet.
            months.append({"sheet": ws.title, "skipped": True})
            continue
        columns = _day_columns(ws, header_row, int(year_match.group(1)))
        if not columns:
            months.append({"sheet": ws.title, "skipped": True})
            warnings.append({"code": "noDayColumns", "sheet": ws.title})
            continue

        sheet_dates: set[str] = set()
        sheet_metrics: set[str] = set()
        backlog_dates: set[str] = set()
        value_count = 0
        for row in range(header_row + 1, ws.max_row + 1):
            label = ws.cell(row, 1).value
            if not isinstance(label, str) or not label.strip():
                # A row without a label but with text in the day columns is
                # the "Heute" summary block's header - stop there.
                if any(isinstance(ws.cell(row, col).value, str) for col in columns):
                    break
                continue
            if label.strip().lower() == "heute":
                break
            key = metric_key(label)
            if not key:
                continue
            for col, day in columns.items():
                value = ws.cell(row, col).value
                if not _is_number(value):
                    continue
                iso = day.isoformat()
                if (iso, key) in values and values[(iso, key)] != float(value):
                    warnings.append({"code": "duplicateDay", "sheet": ws.title, "date": iso, "metric": key})
                values[(iso, key)] = float(value)
                value_count += 1
                sheet_dates.add(iso)
                sheet_metrics.add(key)
                if key in BACKLOG_METRICS:
                    backlog_dates.add(iso)
            if key not in metrics:
                metrics[key] = {"label": label.strip(), "order": len(metrics)}

        first_day = min(columns.values())
        months.append(
            {
                "sheet": ws.title,
                "month": f"{first_day.year}-{first_day.month:02d}",
                "skipped": False,
                "days": len(sheet_dates),
                "metrics": len(sheet_metrics),
                "values": value_count,
                "backlogDays": len(backlog_dates),
                "firstDate": min(sheet_dates) if sheet_dates else None,
                "lastDate": max(sheet_dates) if sheet_dates else None,
            }
        )
        if sheet_dates and not backlog_dates:
            warnings.append({"code": "noBacklog", "sheet": ws.title})

    warnings.extend(_check_backlog_total(values))
    if not values:
        warnings.append({"code": "noValues"})
    return {"values": values, "metrics": metrics, "months": months, "warnings": warnings}


def _check_backlog_total(values: dict) -> list[dict]:
    """The file's "Summe Auftragsbestand" should equal KTL + Pulver +
    Schäumen; list the days where it doesn't (the dashboard uses the sum)."""
    out = []
    days = sorted({day for day, metric in values if metric == BACKLOG_TOTAL_METRIC})
    for day in days:
        parts = [values.get((day, metric)) for metric in BACKLOG_METRICS]
        if any(part is None for part in parts):
            continue
        if abs(sum(parts) - values[(day, BACKLOG_TOTAL_METRIC)]) > 0.01:
            out.append({"code": "totalMismatch", "date": day})
    return out


def carry_forward_days(values: dict) -> set[str]:
    """Saturdays and Sundays whose three backlog values are identical to the
    previous day's - the workbook copies Friday's backlog into the weekend.
    A weekday that happens to repeat the day before (e.g. a Monday equal to
    the weekend) is still a real working day and is not marked."""
    by_day: dict[str, tuple] = {}
    for (day, metric), value in values.items():
        if metric in BACKLOG_METRICS:
            by_day.setdefault(day, {})[metric] = value
    out = set()
    for day, parts in by_day.items():
        if len(parts) != len(BACKLOG_METRICS) or date.fromisoformat(day).weekday() < 5:
            continue
        previous = by_day.get(date.fromordinal(date.fromisoformat(day).toordinal() - 1).isoformat())
        if previous and all(previous.get(m) == parts[m] for m in BACKLOG_METRICS):
            out.add(day)
    return out
